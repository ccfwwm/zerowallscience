"""Offline functional acceptance tests, executed by the release interpreter."""
import argparse
from concurrent.futures import ThreadPoolExecutor
import importlib.metadata as metadata
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import time

ROOT = Path(__file__).resolve().parents[2]
MODULES = '''mcp numpy pandas httpx requests openpyxl pypdf fitz docx pptx matplotlib Bio anndata scanpy mygene gseapy polars pyarrow markitdown liteparse cv2 imagehash pikepdf structlog skimage pydantic_settings markdown yaml imageio tifffile imagecodecs pillow_heif pydicom pylibjpeg jpeg_ls nibabel nrrd SimpleITK astropy dask distributed xarray netCDF4 geopandas pyogrio pyproj shapely pint uncertainties pingouin lifelines pyDOE3 simpy shap pymoo sksurv hdbscan xgboost pymc arviz igraph leidenalg pydeseq2 skbio gprofiler flowio rdkit datamol medchem bids openpiv pyopenms pymatgen.core mp_api.client pytximport qutip qiskit pyzotero cloudscraper playwright.sync_api Crypto.Cipher.AES tomli scipy sklearn statsmodels networkx sympy reportlab pdfplumber sqlalchemy neo4j zarr bs4 lxml umap PIL'''.split()

CASES = {
    'image-comparison': '''
import numpy as np, cv2, imagehash, tifffile, imageio.v3 as iio
from PIL import Image
from skimage.metrics import structural_similarity
from pillow_heif import register_heif_opener
import imagecodecs
a=np.random.default_rng(17).integers(0,255,(128,128),dtype=np.uint8)
assert len(cv2.SIFT_create().detect(a,None)) > 0
im=Image.fromarray(a); assert imagehash.phash(im)-imagehash.phash(im.copy())==0
assert structural_similarity(a,a,data_range=255)==1
tifffile.imwrite('sample.tiff',a); assert np.array_equal(tifffile.imread('sample.tiff'),a)
assert np.array_equal(imagecodecs.png_decode(imagecodecs.png_encode(a)),a)
iio.imwrite('sample.png',a); assert np.array_equal(iio.imread('sample.png'),a)
register_heif_opener(); im.convert('RGB').save('sample.heic'); assert Image.open('sample.heic').size==(128,128)
''',
    'pdf-and-logging': '''
import fitz,pikepdf,markdown,structlog
with fitz.open() as d:
 p=d.new_page();p.insert_text((30,40),'ZeroWall PDF test');d.save('sample.pdf')
with pikepdf.open('sample.pdf') as d: assert len(d.pages)==1
assert '<h1>' in markdown.markdown('# Test')
structlog.get_logger().info('verified',component='release')
''',
    'medical-imaging': '''
import numpy as np,nibabel as nib,nrrd,SimpleITK as sitk,pydicom,jpeg_ls
from pydicom.dataset import FileDataset,FileMetaDataset
from pydicom.uid import ExplicitVRLittleEndian,SecondaryCaptureImageStorage,generate_uid
a=np.arange(64,dtype=np.uint16).reshape(8,8)
meta=FileMetaDataset();meta.TransferSyntaxUID=ExplicitVRLittleEndian;meta.MediaStorageSOPClassUID=SecondaryCaptureImageStorage;meta.MediaStorageSOPInstanceUID=generate_uid()
d=FileDataset('test.dcm',{},file_meta=meta,preamble=b'\\0'*128)
d.SOPClassUID=meta.MediaStorageSOPClassUID;d.SOPInstanceUID=meta.MediaStorageSOPInstanceUID
d.Rows,d.Columns=a.shape;d.SamplesPerPixel=1;d.PhotometricInterpretation='MONOCHROME2';d.BitsAllocated=16;d.BitsStored=16;d.HighBit=15;d.PixelRepresentation=0;d.PixelData=a.tobytes();d.save_as('test.dcm',enforce_file_format=True)
assert np.array_equal(pydicom.dcmread('test.dcm').pixel_array,a)
assert np.array_equal(jpeg_ls.decode(jpeg_ls.encode(a)),a)
for uid in ['1.2.840.10008.1.2.4.50','1.2.840.10008.1.2.4.90','1.2.840.10008.1.2.5']:
 assert pydicom.pixels.get_decoder(uid).is_available,uid
volume=np.arange(64,dtype=np.float32).reshape(4,4,4)
nib.save(nib.Nifti1Image(volume,np.eye(4)),'test.nii');assert np.array_equal(nib.load('test.nii').get_fdata(),volume)
nrrd.write('test.nrrd',volume);assert np.array_equal(nrrd.read('test.nrrd')[0],volume)
image=sitk.GetImageFromArray(volume);sitk.WriteImage(image,'test.mha');assert np.array_equal(sitk.GetArrayFromImage(sitk.ReadImage('test.mha')),volume)
''',
    'zotero-mocked-api': '''
from pyzotero import zotero
import httpx
z=zotero.Zotero('12345','user','offline-test-key')
def reply(request):
 assert 'api.zotero.org' in str(request.url)
 return httpx.Response(200,headers={'Content-Type':'application/json','Total-Results':'1','Last-Modified-Version':'1'},json=[{'key':'ABC123','version':1,'data':{'itemType':'journalArticle','title':'Offline release test'}}],request=request)
z.client.close();z.client=httpx.Client(transport=httpx.MockTransport(reply))
items=z.items();assert items[0]['data']['title']=='Offline release test';z.client.close()
''',
    'office-conversions': '''
from markitdown import MarkItDown
from docx import Document
from pptx import Presentation
from pptx.util import Inches
import openpyxl,fitz
doc=Document();doc.add_paragraph('ZeroWall Office marker');doc.save('sample.docx')
p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text='ZeroWall Office marker';p.save('sample.pptx')
w=openpyxl.Workbook();w.active.append(['ZeroWall Office marker',42]);w.save('sample.xlsx')
with fitz.open() as d:
 page=d.new_page();page.insert_text((30,40),'ZeroWall Office marker');d.save('sample.pdf')
converter=MarkItDown()
for path in ['sample.docx','sample.pptx','sample.xlsx','sample.pdf']:
 result=converter.convert(path);assert 'ZeroWall' in result.text_content,(path,result.text_content)
''',
    'statistics': '''
import numpy as np,pandas as pd,pingouin as pg,pint,uncertainties,pyDOE3,simpy
from lifelines import KaplanMeierFitter
assert pg.ttest([1,2,3,4,5],[2,3,4,5,6]).shape[0]==1
assert KaplanMeierFitter().fit([1,2,3],[1,1,0]).survival_function_.shape[0]>0
assert (1*pint.UnitRegistry().meter).to('centimeter').magnitude==100
assert uncertainties.ufloat(2,.1).nominal_value==2
assert pyDOE3.ff2n(3).shape==(8,3)
e=simpy.Environment();e.run(until=2);assert e.now==2
''',
    'ml-and-optimization': '''
import numpy as np,shap,hdbscan,xgboost
from sklearn.ensemble import RandomForestRegressor
from sksurv.nonparametric import kaplan_meier_estimator
from pymoo.problems import get_problem
x=np.random.default_rng(1).normal(size=(30,3));y=x[:,0]*2
m=RandomForestRegressor(n_estimators=3,random_state=1).fit(x,y)
assert shap.TreeExplainer(m)(x[:2]).values.shape==(2,3)
assert len(hdbscan.HDBSCAN(min_cluster_size=3).fit_predict(x))==30
assert len(xgboost.XGBRegressor(n_estimators=2,n_jobs=1).fit(x,y).predict(x))==30
assert get_problem('sphere').evaluate(np.zeros((1,10)))[0]==0
t,s=kaplan_meier_estimator(np.array([True,True,False]),np.array([1,2,3]));assert len(t)==3
''',
    'bayesian': '''
import pymc as pm,arviz as az,numpy as np
with pm.Model() as m:
 x=pm.Normal('x',0,1);pm.Normal('obs',x,1,observed=np.array([0.1,0.2]))
 result=pm.sample_prior_predictive(samples=3,random_seed=1)
assert result.prior.sizes['draw']==3
assert az.from_dict(posterior={'a':np.ones((1,10))}).posterior.sizes['draw']==10
''',
    'data-and-geography': '''
import numpy as np,dask.array as da,xarray as xr,geopandas as gpd,pyogrio
from distributed import Client
from shapely.geometry import Point
from pyproj import Transformer
from astropy import units as u
assert da.arange(10,chunks=3).sum().compute()==45
with Client(processes=False,n_workers=1,threads_per_worker=1,dashboard_address=None) as client: assert client.submit(sum,[1,2,3]).result()==6
x=xr.Dataset({'a':('n',np.arange(5))});x.to_netcdf('data.nc');assert xr.load_dataset('data.nc').a.sum()==10
g=gpd.GeoDataFrame({'a':[1]},geometry=[Point(0,0)],crs=4326);g.to_file('geo.gpkg',driver='GPKG',engine='pyogrio');assert len(pyogrio.read_dataframe('geo.gpkg'))==1
assert Transformer.from_crs(4326,3857,always_xy=True).transform(0,0)==(0,0)
assert (1*u.km).to(u.m).value==1000
''',
    'single-cell': '''
import numpy as np,scanpy as sc,anndata as ad
x=ad.AnnData(np.random.default_rng(2).poisson(5,(40,20)).astype(float))
sc.pp.normalize_total(x);sc.pp.log1p(x);sc.pp.pca(x,n_comps=8)
sc.pp.neighbors(x,n_neighbors=5,n_pcs=8);sc.tl.leiden(x,flavor='igraph',n_iterations=2,directed=False)
assert 'leiden' in x.obs and len(x.obs)==40
''',
    'deseq-and-bio': '''
import numpy as np,pandas as pd,skbio
from pydeseq2.dds import DeseqDataSet
from pydeseq2.ds import DeseqStats
from gprofiler import GProfiler
counts=pd.DataFrame(np.random.default_rng(7).negative_binomial(10,.5,(8,30)),index=[f's{i}' for i in range(8)])
meta=pd.DataFrame({'condition':['A']*4+['B']*4},index=counts.index)
d=DeseqDataSet(counts=counts,metadata=meta,design='~condition',n_cpus=1,quiet=True);d.deseq2()
s=DeseqStats(d,contrast=['condition','B','A'],n_cpus=1,quiet=True);s.summary();assert len(s.results_df)==30
assert str(skbio.DNA('ACGT').reverse_complement())=='ACGT'
assert GProfiler(return_dataframe=True).return_dataframe
''',
    'chemistry-materials': '''
from rdkit import Chem
import datamol as dm,medchem as mc
from pymatgen.core import Lattice,Structure
from mp_api.client import MPRester
assert Chem.MolToSmiles(Chem.MolFromSmiles('CCO'))=='CCO'
assert dm.to_mol('CCO').GetNumAtoms()==3
assert len(mc.functional.list_default_available_rules())>0
s=Structure(Lattice.cubic(3.5),['Na','Cl'],[[0,0,0],[.5,.5,.5]]);assert len(Structure.from_str(s.to(fmt='cif'),fmt='cif'))==2
with MPRester(api_key='0'*32,mute_progress_bars=True) as client: assert client is not None
''',
    'quantum-and-massspec': '''
import qutip as qt
from qiskit import QuantumCircuit
from qiskit.quantum_info import Statevector
import pyopenms as oms
assert abs(qt.basis(2,0).norm()-1)<1e-12
c=QuantumCircuit(1);c.h(0);assert abs(Statevector.from_instruction(c).probabilities()[0]-.5)<1e-12
s=oms.MSSpectrum();s.set_peaks(([100.,200.],[1.,2.]));assert s.size()==2
''',
    'flowio-bids-piv': '''
import flowio,numpy as np,json
from bids import BIDSLayout
from openpiv import pyprocess
flowio.create_fcs('sample.fcs',list(map(float,range(12))),['A','B','C'])
f=flowio.FlowData('sample.fcs');assert f.event_count==4
from pathlib import Path
Path('bids').mkdir(exist_ok=True);Path('bids/dataset_description.json').write_text(json.dumps({'Name':'test','BIDSVersion':'1.10.0'}))
assert BIDSLayout('bids',validate=False).description['Name']=='test'
a=np.random.default_rng(5).integers(0,255,(64,64),dtype=np.int32)
u,v,s=pyprocess.extended_search_area_piv(a,a,window_size=32,overlap=16,dt=1,search_area_size=32)
assert np.nanmax(abs(u))<.1 and np.nanmax(abs(v))<.1
''',
    'pytximport': '''
import inspect,pytximport
assert callable(pytximport.tximport)
from pathlib import Path
for name in ['s1','s2']:
 Path(name).mkdir(exist_ok=True);Path(name+'/quant.sf').write_text('Name\\tLength\\tEffectiveLength\\tTPM\\tNumReads\\nt1\\t1000\\t800\\t600000\\t60\\nt2\\t1000\\t800\\t400000\\t40\\n')
result=pytximport.tximport(['s1/quant.sf','s2/quant.sf'],data_type='salmon',tx_out=True)
assert result is not None
''',
    'user-tools': '''
from playwright.sync_api import sync_playwright
from Crypto.Cipher import AES
import tomli,cloudscraper
with sync_playwright() as p: assert p.chromium.name=='chromium'
c=AES.new(b'0'*16,AES.MODE_ECB);assert len(c.encrypt(b'1'*16))==16
assert tomli.loads('a=1')['a']==1
s=cloudscraper.create_scraper();assert s is not None;s.close()
''',
}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', required=True)
    parser.add_argument('--case')
    args = parser.parse_args()
    output = Path(args.output).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    work = output.parent / 'functional'
    work.mkdir(exist_ok=True)
    env = {**os.environ, 'PYTHONNOUSERSITE': '1', 'PYTHONPATH': '', 'MPLBACKEND': 'Agg', 'NUMBA_NUM_THREADS': '2', 'OMP_NUM_THREADS': '2', 'OPENBLAS_NUM_THREADS': '2'}
    if sys.version_info[:3] != (3,12,10):
        raise SystemExit('Expected CPython 3.12.10')
    site = Path(sys.executable).parent / 'site-packages'
    if not sys.flags.no_user_site:
        raise SystemExit('Run verification with -s')
    packages = {re.sub(r'[-_.]+','-',d.metadata['Name']).lower():d.version for d in metadata.distributions(path=[str(site)])}
    def child(name, code, timeout=180):
        start=time.monotonic(); cwd=work/name;cwd.mkdir(exist_ok=True)
        try:
            r=subprocess.run([sys.executable,'-s','-B','-c',code],cwd=cwd,env=env,capture_output=True,text=True,encoding='utf-8',errors='replace',timeout=timeout)
            result={'ok':r.returncode==0,'seconds':round(time.monotonic()-start,2),'output':(r.stdout+'\n'+r.stderr)[-5000:]}
        except subprocess.TimeoutExpired:
            result={'ok':False,'seconds':timeout,'output':'Timed out'}
        print(name, 'PASS' if result['ok'] else 'FAIL', flush=True)
        return result
    report={'python':'3.12.10','isolated':True,'packages':packages,'imports':{},'cases':{}}
    if not args.case:
        def probe(module):
            return module, child('import-'+module, f'import importlib,pathlib; m=importlib.import_module({module!r}); p=getattr(m,"__file__",None); assert p is None or pathlib.Path(p).resolve().is_relative_to(pathlib.Path({str(site)!r}).resolve()), p; print(p)')
        with ThreadPoolExecutor(max_workers=4) as pool:
            for name,result in pool.map(probe,MODULES): report['imports'][name]=result
    elif output.exists():
        report=json.loads(output.read_text('utf-8'))
    for name,code in CASES.items():
        if not args.case or name==args.case:
            report['cases'][name]=child(name,code,240)
            output.write_text(json.dumps(report,indent=2),encoding='utf-8')
    report['ok']=all(r['ok'] for r in [*report['imports'].values(),*report['cases'].values()])
    output.write_text(json.dumps(report,indent=2),encoding='utf-8')
    print('Report:',output,flush=True)
    raise SystemExit(0 if report['ok'] else 1)


if __name__=='__main__':main()
