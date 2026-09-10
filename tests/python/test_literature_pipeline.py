import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch


SCRIPT = Path(__file__).resolve().parents[2] / "resources" / "skills" / "zerowall-literature" / "scripts" / "literature_pipeline.py"
spec = importlib.util.spec_from_file_location("literature_pipeline", SCRIPT)
module = importlib.util.module_from_spec(spec)
assert spec and spec.loader
sys.modules[spec.name] = module
spec.loader.exec_module(module)


class LiteraturePipelineTests(unittest.TestCase):
    def test_title_score_prefers_exact_match(self):
        exact = module.title_score("A study of WDR6 in mice", "A study of WDR6 in mice")
        distant = module.title_score("A study of WDR6 in mice", "Genome-wide association of bone density")
        self.assertGreater(exact, distant)
        self.assertEqual(exact, 1.0)

    def test_citation_contexts_capture_multiple_locations(self):
        text = "Results support the finding [18]. Later, unlike the prior work [18], the mechanism differed."
        target = module.Paper(key="pmid:1", title="Prior work", authors=["Alice Smith"])
        contexts = module.citation_contexts(text, target, ref_number=18)
        self.assertGreaterEqual(len(contexts), 2)
        self.assertTrue(all(item["excerpt"] for item in contexts))

    def test_citation_contexts_expand_numeric_ranges(self):
        target = module.Paper(key="doi:10.1/x", title="Target", authors=["Alice Smith"])
        contexts = module.citation_contexts("The result was cited by refs [17-20] and again (18, 22).", target, ref_number=18)
        self.assertEqual(len(contexts), 2)
        self.assertTrue(all(item["page"] == 1 for item in contexts))

    def test_reference_number_is_inferred_from_bibliography(self):
        text = "Introduction cites the work [17-20].\nReferences\n18. Smith A. Upregulation of WDR6 drives hepatic de novo lipogenesis in insulin resistance in mice. doi:10.1038/s42255-023-00896-7"
        target = module.Paper(key="doi:10.1038/s42255-023-00896-7", title="Upregulation of WDR6 drives hepatic de novo lipogenesis in insulin resistance in mice", doi="10.1038/s42255-023-00896-7")
        self.assertEqual(module.infer_reference_number(text, target), 18)
        contexts = module.citation_contexts(text, target, 18)
        self.assertEqual(len(contexts), 1)

    def test_dedup_merges_doi_and_title_aliases(self):
        first = module.Paper(key="openalex:W1", title="A WDR6 study", doi="10.1000/ABC", year="2023", authors=["A"])
        second = module.Paper(key="crossref:1", title="A WDR6 study", doi="https://doi.org/10.1000/abc", year="2023", authors=["B"], countries=["US"])
        papers, aliases = module.deduplicate_papers([first, second])
        self.assertEqual(len(papers), 1)
        self.assertEqual(len(aliases), 1)
        self.assertIn("B", papers[0].authors)
        self.assertIn("US", papers[0].countries)

    def test_pdf_validation_requires_magic_and_returns_hash(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "paper.pdf"
            path.write_bytes(b"%PDF-1.7\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF\n")
            paper = module.Paper(key="doi:10.1/test", title="A paper")
            ok, reason, digest = module.validate_pdf(path, paper)
            self.assertTrue(ok)
            self.assertEqual(reason, "ok")
            self.assertEqual(len(digest), 64)

    def test_task_writes_json_state_atomically(self):
        with tempfile.TemporaryDirectory() as directory:
            task = module.Task(Path(directory))
            state = task.load()
            state["stage"] = "identified"
            task.save(state)
            self.assertEqual(task.load()["stage"], "identified")
            self.assertFalse(task.state_path.with_suffix(".tmp").exists())

    def test_local_csv_rejects_multiple_main_article_titles(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "articles.csv"
            path.write_text("title\nFirst article title with enough words to be detected\nSecond article title with enough words to be detected\n", encoding="utf-8")
            client = module.Client(module.Task(Path(directory) / "task"))
            with self.assertRaises(ValueError):
                client.from_local(path)

    def test_local_csv_accepts_one_main_article_title(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "article.csv"
            title = "Upregulation of WDR6 drives hepatic de novo lipogenesis in insulin resistance in mice"
            path.write_text(f"title\n{title}\n", encoding="utf-8")
            client = module.Client(module.Task(Path(directory) / "task"))
            paper = client.from_local(path)
            self.assertEqual(paper.title, title)

    def test_flatten_excel_value_never_returns_nested_values(self):
        value = module.flatten_excel_value([{"a": [1, 2]}, "x"])
        self.assertIsInstance(value, str)
        self.assertIn("a:", value)

    def test_analyze_defaults_to_full_pdf_acquisition(self):
        captured = {}

        def capture(args):
            captured["download_pdfs"] = args.download_pdfs
            captured["directions"] = args.directions
            captured["max_papers"] = args.max_papers
            return 0

        with patch.object(module, "run_analyze", capture):
            result = module.main(["analyze", "A paper title", "--output", "task"])
        self.assertEqual(result, 0)
        self.assertTrue(captured["download_pdfs"])
        self.assertEqual(captured["directions"], "cited-by")
        self.assertIsNone(captured["max_papers"])

    def test_ingest_mineru_result_persists_artifacts_and_provenance(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "task"
            task = module.Task(root)
            paper = module.Paper(key="doi:10.1/test", title="Test paper", doi="10.1/test", pdf_path=str(root / "downloads" / "test.pdf"))
            state = task.load()
            state.update({"task_root": str(root.resolve()), "papers": [module.asdict(paper)]})
            task.save(state)
            run_dir = Path(directory) / "mineru-run"
            run_dir.mkdir()
            (run_dir / "full.md").write_text("# Test paper\n\nMinerU extracted body text.", encoding="utf-8")
            (run_dir / "run.json").write_text("{}", encoding="utf-8")

            with patch.object(module, "report"):
                parsed = module.ingest_mineru_result(task, state, "10.1/test", run_dir, "task-123", "precision")

            self.assertEqual(parsed.parse_status, "mineru_parsed")
            self.assertEqual(parsed.parser_task_id, "task-123")
            self.assertEqual(parsed.parser_api, "precision")
            self.assertTrue(Path(parsed.parsed_text_path).is_file())
            self.assertEqual(len(parsed.parser_artifacts), 2)
            self.assertEqual(task.load()["stage"], "analysis_pending")
            self.assertTrue(any(row["provider"] == "mineru" for row in task.ledger()))

    def test_ingest_mineru_result_preserves_nested_assets_and_detects_broken_links(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "task"
            task = module.Task(root)
            paper = module.Paper(key="doi:10.1/assets", title="Asset paper", doi="10.1/assets", pdf_path="paper.pdf")
            state = task.load(); state.update({"task_root": str(root.resolve()), "papers": [module.asdict(paper)]}); task.save(state)
            run_dir = Path(directory) / "mineru-run"; (run_dir / "images").mkdir(parents=True); (run_dir / "tables").mkdir()
            (run_dir / "full.md").write_text("# Paper\n\n![figure](images/figure.png)\n![missing](images/missing.png)", encoding="utf-8")
            (run_dir / "images" / "figure.png").write_bytes(b"png")
            (run_dir / "tables" / "table.json").write_text("{}", encoding="utf-8")
            with patch.object(module, "report"):
                parsed = module.ingest_mineru_result(task, state, "10.1/assets", run_dir, "run-1", "mineru")
            self.assertTrue((Path(parsed.mineru_snapshot_dir) / "images" / "figure.png").is_file())
            self.assertTrue((Path(parsed.mineru_snapshot_dir) / "tables" / "table.json").is_file())
            self.assertIn("images/missing.png", parsed.mineru_broken_links)
            self.assertGreaterEqual(parsed.mineru_file_counts["total"], 3)

    def test_paper_text_requires_mineru_snapshot(self):
        paper = module.Paper(key="doi:10.1/text", title="Paper", pdf_path="missing.pdf")
        self.assertEqual(module.paper_text(paper), "")

    def test_local_pdf_is_registered_without_text_extraction(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "opaque-name.pdf"
            source.write_bytes(b"%PDF-1.4\n%%EOF\n")
            task = module.Task(Path(directory) / "task")
            paper = module.Client(task).from_local(source)
            with patch.object(module, "extract_pdf_text", side_effect=AssertionError("must not parse local PDF")):
                module.acquire_provided_target_pdf(task, paper, source)
            self.assertEqual(paper.pdf_status, "provided_pdf")
            self.assertEqual(paper.parse_status, "mineru_required")
            self.assertTrue(Path(paper.pdf_path).is_file())

    def test_pdf_paths_are_flat_and_reader_facing(self):
        with tempfile.TemporaryDirectory() as directory:
            task = module.Task(Path(directory) / "task")
            target = module.Paper(key="doi:target", title="Target title", direction="target")
            cited = module.Paper(key="doi:cited", title="Citing title", direction="cited-by", authors=["Ada Lovelace"], year="2025")
            target_path = module.canonical_pdf_path(task, target)
            cited_path = module.canonical_pdf_path(task, cited)
            self.assertEqual(target_path.parent, task.root / "downloads")
            self.assertEqual(cited_path.parent, task.root / "downloads")
            self.assertTrue(target_path.name.startswith("target__"))
            self.assertTrue(cited_path.name.startswith("cited__Ada Lovelace__2025__"))

    def test_simplified_ingest_rejects_cited_paper(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "task"; task = module.Task(root)
            target = module.Paper(key="doi:10.1/target", title="Target", direction="target", pdf_path="target.pdf")
            cited = module.Paper(key="doi:10.1/cited", title="Cited", direction="cited-by", pdf_path="cited.pdf")
            state = task.load(); state.update({"workflow_mode": module.WORKFLOW_MODE, "task_root": str(root), "papers": [module.asdict(target), module.asdict(cited)]}); task.save(state)
            run_dir = Path(directory) / "run"; run_dir.mkdir(); (run_dir / "full.md").write_text("# Cited paper\n\nBody", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "target paper"):
                module.ingest_mineru_result(task, state, cited.key, run_dir)

    def test_simplified_target_ingest_unlocks_cited_by(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "task"; task = module.Task(root)
            target = module.Paper(key="local:test", title="file-name", direction="target", source="local", pdf_path="target.pdf", pdf_status="provided_pdf", acquisition_state="acquisition_terminal")
            state = task.load(); state.update({"workflow_mode": module.WORKFLOW_MODE, "task_root": str(root), "papers": [module.asdict(target)]}); task.save(state)
            run_dir = Path(directory) / "run"; run_dir.mkdir(); (run_dir / "full.md").write_text("# A precise target paper title\n\nBody", encoding="utf-8")
            with patch.object(module, "report"):
                parsed = module.ingest_mineru_result(task, state, target.key, run_dir, "task-1")
            saved = task.load()
            self.assertEqual(parsed.title, "A precise target paper title")
            self.assertEqual(saved["stage"], "target_mineru_parsed")
            self.assertEqual(saved["phase_status"]["cited_by_expanded"], "pending")

    def test_citation_relation_never_contains_body_markers(self):
        target = module.Paper(key="doi:target", title="Target", authors=["Alice Smith"])
        cited = module.Paper(key="doi:cited", title="Follow-up", authors=["Bob Jones"], direction="cited-by")
        relation = module.citation_relation(target, cited)
        for forbidden in ("marker", "offset", "page", "excerpt"):
            self.assertNotIn(forbidden, relation)

    def test_simplified_report_has_explanation_sheet_and_no_remote_runtime(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "task"; task = module.Task(root)
            target = module.Paper(key="doi:target", title="Target paper", direction="target", authors=["Alice Smith"], pdf_path="target.pdf", pdf_status="provided_pdf", acquisition_state="acquisition_terminal", parse_status="mineru_parsed", parsed_text_path="full.md")
            cited = module.Paper(key="doi:cited", title="Citing paper", direction="cited-by", pdf_status="unavailable_no_authorized_source", acquisition_state="acquisition_terminal", parse_status="not_required", authors=["Alice Smith", "Bob Jones"], author_records=[{"name": "Alice Smith", "current_title": "Professor", "sources": ["https://example.org/alice"]}, {"name": "Bob Jones", "current_title": "Professor", "current_institution": "Example University", "sources": ["https://example.org/bob"]}], citation_relation=module.citation_relation(target, module.Paper(key="doi:cited", title="Citing paper", direction="cited-by")))
            state = task.load(); state.update({"workflow_mode": module.WORKFLOW_MODE, "task_root": str(root), "stage": "report_building", "papers": [module.asdict(target), module.asdict(cited)]})
            module.simplified_report(task, state, [target, cited])
            from openpyxl import load_workbook
            book = load_workbook(root / "papers.xlsx", read_only=True)
            try:
                self.assertEqual(book.sheetnames, ["引文列表", "作者列表"])
                forbidden = {"PDF来源", "PDF路径", "PDF获取情况", "SHA256", "MinerU状态", "MinerU正文", "置信度", "身份状态"}
                for sheet in book.worksheets:
                    headers = {cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))}
                    self.assertTrue(forbidden.isdisjoint(headers))
                citation_headers = {cell.value for cell in next(book["引文列表"].iter_rows(min_row=1, max_row=1))}
                self.assertIn("PDF文件名", citation_headers)
                workbook_text = " ".join(str(cell.value or "") for sheet in book.worksheets for row in sheet.iter_rows() for cell in row)
                self.assertNotIn("Alice Smith", workbook_text)
                self.assertIn("Bob Jones", workbook_text)
            finally:
                book.close()
            html = (root / "report.html").read_text(encoding="utf-8")
            self.assertNotIn("cdn", html.lower())
            self.assertNotIn("<script src=", html.lower())
            for forbidden in ("MinerU", "PDF 覆盖", "PDF获取", "OpenAlex", "Semantic Scholar", "未确认", "not_found"):
                self.assertNotIn(forbidden, html)
            self.assertNotIn("Alice Smith", html)
            self.assertIn("Bob Jones", html)
            self.assertIn("目标论文引文影响与作者画像", html)
            self.assertIn("多维引文统计", html)
            self.assertIn("核心引文作者画像", html)
            self.assertEqual(html.count('<section class="strip"'), 1)
            self.assertEqual(html.count("去重后的 cited-by 引文"), 1)
            self.assertNotIn("目标论文证据解剖", html)
            self.assertFalse((root / "report_academic.html").exists())
            self.assertTrue((root / "analysis" / "report_manifest.json").is_file())

    def test_target_author_is_excluded_even_when_authoring_a_citing_paper(self):
        with tempfile.TemporaryDirectory() as directory:
            task = module.Task(Path(directory) / "task")
            target = module.Paper(key="doi:target", title="Target", direction="target", authors=["Alice Smith"])
            cited = module.Paper(key="doi:cited", title="Self citing follow-up", direction="cited-by", authors=["Alice Smith", "Bob Jones"])
            entities = module.unique_author_entities([target, cited])
            self.assertEqual([row["name"] for row in entities], ["Bob Jones"])
            with patch.dict(module.os.environ, {"LITERATURE_WEB_SEARCH_MODULE": ""}, clear=False):
                records = module.enrich_all_authors(task, [target, cited])
            self.assertEqual([row["name"] for row in records], ["Bob Jones"])

    def test_title_analyze_acquires_only_target_before_mineru(self):
        with tempfile.TemporaryDirectory() as directory:
            target = module.Paper(key="doi:10.1/target", title="A target paper", doi="10.1/target", direction="target")
            def downloaded(_client, paper, _allow_tsg=True):
                paper.pdf_path = str(Path(directory) / "target.pdf")
                paper.pdf_status = "downloaded_open_access"
                paper.acquisition_state = "acquisition_terminal"
            with patch.object(module.Client, "resolve", return_value=(target, [])), patch.object(module.Client, "enrich", side_effect=lambda paper: paper), patch.object(module.Client, "expand_openalex") as expand, patch.object(module, "download_paper", side_effect=downloaded), patch.object(module, "report"):
                result = module.main(["analyze", target.title, "--output", str(Path(directory) / "task")])
            self.assertEqual(result, 0)
            expand.assert_not_called()
            state = module.Task(Path(directory) / "task").load()
            self.assertEqual(state["stage"], "target_mineru_required")
            self.assertEqual(len(state["papers"]), 1)

    def test_local_pdf_analyze_waits_for_target_mineru(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "provided.pdf"; source.write_bytes(b"%PDF-1.4\n%%EOF\n")
            task_root = Path(directory) / "task"
            with patch.object(module.Client, "expand_openalex") as expand, patch.object(module, "report"):
                result = module.main(["analyze", str(source), "--output", str(task_root)])
            self.assertEqual(result, 0); expand.assert_not_called()
            state = module.Task(task_root).load(); paper = state["papers"][0]
            self.assertEqual(state["stage"], "target_mineru_required")
            self.assertEqual(paper["pdf_source"], "user_provided")
            self.assertEqual(paper["parse_status"], "mineru_required")

    def test_resume_does_not_expand_before_target_mineru(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "task"; task = module.Task(root)
            target = module.Paper(key="doi:target", title="Target", direction="target", pdf_path="target.pdf", pdf_status="provided_pdf", acquisition_state="acquisition_terminal", parse_status="mineru_required")
            state = task.load(); state.update({"workflow_mode": module.WORKFLOW_MODE, "task_root": str(root), "single_article": True, "papers": [module.asdict(target)]}); task.save(state)
            with patch.object(module.Client, "expand_openalex") as expand, patch.object(module, "report"):
                result = module.main(["resume", str(root)])
            self.assertEqual(result, 0); expand.assert_not_called()
            self.assertEqual(task.load()["stage"], "target_mineru_required")

    def test_finalize_does_not_require_cited_paper_mineru(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "task"; task = module.Task(root)
            full_md = root / "parsed" / "target" / "run" / "full.md"; full_md.parent.mkdir(parents=True); full_md.write_text("# Target\n\nTarget evidence body.", encoding="utf-8")
            target = module.Paper(key="doi:target", title="Target", direction="target", pdf_path="target.pdf", pdf_status="provided_pdf", acquisition_state="acquisition_terminal", parse_status="mineru_parsed", parsed_text_path=str(full_md))
            cited = module.Paper(key="doi:cited", title="Cited", direction="cited-by", pdf_status="unavailable_no_authorized_source", acquisition_state="acquisition_terminal", parse_status="not_required", authors=["Bob Jones"])
            state = task.load(); state.update({"workflow_mode": module.WORKFLOW_MODE, "task_root": str(root), "single_article": True, "cited_by_expanded": True, "stage": "report_building", "papers": [module.asdict(target), module.asdict(cited)]})
            module.build_simplified_analysis(task, state, [target, cited]); state["papers"] = [module.asdict(target), module.asdict(cited)]; task.save(state); module.report(task, state)
            module.finalize_analysis(task, state)
            self.assertEqual(task.load()["stage"], "partial_complete")

    def test_tsg_is_enabled_by_default_but_can_be_disabled(self):
        with tempfile.TemporaryDirectory() as directory:
            task = module.Task(Path(directory) / "task")
            client = module.Client(task)
            paper = module.Paper(key="doi:10.1/tsg", title="TSG paper", doi="10.1/tsg")
            with patch.object(module, "download_via_tsg", return_value=True) as tsg, patch.object(module, "download_via_paper_download", return_value=False) as paper_download, patch.object(module, "download_authorized_adapter", return_value=False):
                with patch.dict(module.os.environ, {"TSG_PM_JSESSIONID": "x", "TSG_SESSIONID": "x", "TSG_SGUSER": "x", "TSG_TSGUSER": "x"}, clear=False):
                    module.download_paper(client, paper)
            tsg.assert_called_once()
            # TSG is the first authorized fallback.  paper-download is only
            # reached when the title search cannot acquire a verified PDF.
            paper_download.assert_not_called()

    def test_shadow_cascade_can_be_explicitly_disabled(self):
        with tempfile.TemporaryDirectory() as directory:
            task = module.Task(Path(directory) / "task")
            client = module.Client(task)
            paper = module.Paper(key="doi:10.1/no-shadow", title="No shadow paper", doi="10.1/no-shadow")
            with patch.object(module, "download_via_paper_download", return_value=False) as paper_download, patch.object(module, "download_via_tsg", return_value=False), patch.object(module, "download_authorized_adapter", return_value=False):
                with patch.dict(module.os.environ, {"RESEARCH_ENABLE_SHADOW_LIBS": "0"}, clear=False):
                    module.download_paper(client, paper)
            self.assertFalse(paper_download.call_args.kwargs["allow_shadow"])

    def test_tsg_credential_accepts_packaged_environment_and_cookie_export(self):
        with patch.dict(module.os.environ, {"tsg_pm_jsessionid": "JSESSIONID=pm-value"}, clear=True):
            value, source = module._tsg_credential("TSG_PM_JSESSIONID")
            self.assertEqual(value, "pm-value")
            self.assertEqual(source, "environment")
        exported = json.dumps({"cookies": [
            {"domain": ".yuntsg.com", "name": "sguser", "value": "sg-value"},
            {"domain": "pm.yuntsg.com", "name": "JSESSIONID", "value": "pm-exported"},
        ]})
        with patch.dict(module.os.environ, {"RESEARCH_BROWSER_COOKIES": exported}, clear=True):
            self.assertEqual(module._tsg_credential("TSG_PM_JSESSIONID"), ("pm-exported", "browser_cookie_export"))
            self.assertEqual(module._tsg_credential("TSG_SGUSER"), ("sg-value", "browser_cookie_export"))

    def test_missing_tsg_credential_is_a_blocked_search_attempt_not_a_silent_skip(self):
        with tempfile.TemporaryDirectory() as directory:
            task = module.Task(Path(directory) / "task")
            client = module.Client(task)
            paper = module.Paper(key="doi:10.1/tsg", title="Exact TSG search title", doi="10.1/tsg")
            with patch.dict(module.os.environ, {}, clear=True):
                self.assertFalse(module.download_via_tsg(client, paper, Path(directory) / "paper.pdf"))
            attempt = paper.acquisition_attempts[-1]
            self.assertEqual(attempt["status"], "blocked_missing_credentials")
            self.assertEqual(attempt["query"], paper.title)
            self.assertIn("TSG_PM_JSESSIONID", attempt["missing"])
            self.assertNotIn("skipped", attempt["status"])

    def test_finalize_rejects_missing_agent_analysis(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "task"
            task = module.Task(root)
            paper = module.Paper(key="doi:10.1/test", title="Test paper", doi="10.1/test", direction="target", pdf_path="paper.pdf", parse_status="mineru_parsed")
            state = task.load()
            state.update({"task_root": str(root.resolve()), "papers": [module.asdict(paper)]})
            with self.assertRaisesRegex(ValueError, "analyze_authors"):
                module.finalize_analysis(task, state)
            self.assertEqual(task.load()["stage"], "analysis_pending")

    def test_finalize_marks_complete_when_required_evidence_exists(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "task"
            task = module.Task(root)
            analysis = root / "analysis"
            analysis.mkdir(exist_ok=True)
            (analysis / "provider_evidence.json").write_text('{"queries":[{"provider":"OpenAlex","status":"ok"}]}', encoding="utf-8")
            for name in ("citation_analysis.md", "author_analysis.md", "synthesis.md"):
                (analysis / name).write_text(f"# {name}\n\n" + "Verified evidence with source identifiers and explicit limitations. " * 2, encoding="utf-8")
            paper = module.Paper(key="doi:10.1/test", title="Test paper", doi="10.1/test", direction="target", pdf_path="paper.pdf", parse_status="mineru_parsed")
            state = task.load()
            state.update({"task_root": str(root.resolve()), "papers": [module.asdict(paper)]})
            with patch.object(module, "report"):
                module.finalize_analysis(task, state)
            self.assertEqual(task.load()["stage"], "complete")


if __name__ == "__main__":
    unittest.main()
