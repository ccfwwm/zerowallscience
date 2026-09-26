"""Focused 7.1.0 integrity contracts; runs in the managed Python environment."""
from __future__ import annotations

import json
import os
from pathlib import Path
import sys
import tempfile
import time
import unittest
import weakref
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / "resources/skills/zerowall-image-dup/scripts"
sys.path[:0] = [str(SCRIPTS), str(SCRIPTS / "vendor")]

import cv2
import numpy as np
from PIL import Image
from manusift.contracts import ExtractedImage, ParsedDoc
from manusift.detectors.image_forensics import _panel_then_match_findings, _replace_checkpoint
from manusift.detectors.sift_copymove import CrossMatchAnalysis, match_two_arrays
from numeric_audit import _calculate_formula, _xlsx_rows, audit
from local_ocr import scan, scan_isolated
from zerowall_integrity import coverage_from_steps, save


class Integrity710(unittest.TestCase):
    def test_atomic_progress_save_retries_windows_reader(self):
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "progress.json"
            target.write_text("{}", encoding="utf-8")
            replace = Path.replace
            attempts = 0

            def busy_once(source, destination):
                nonlocal attempts
                attempts += 1
                if attempts == 1:
                    raise PermissionError("reader briefly holds the destination")
                return replace(source, destination)

            with patch.object(Path, "replace", busy_once):
                save(target, {"completed": 3})
            self.assertEqual(attempts, 2)
            self.assertEqual(json.loads(target.read_text(encoding="utf-8")), {"completed": 3})
            temp = Path(folder) / "pair.tmp"
            temp.write_text('{"compared": 1}', encoding="utf-8")
            attempts = 0
            with patch.object(Path, "replace", busy_once):
                _replace_checkpoint(temp, target)
            self.assertEqual(attempts, 2)
            self.assertEqual(json.loads(target.read_text(encoding="utf-8")), {"compared": 1})

    def test_flip_vertical_and_blank_gate(self):
        original = np.random.default_rng(17).integers(0, 255, (280, 320, 3), dtype=np.uint8)
        match = match_two_arrays(original, cv2.flip(original, 0))
        self.assertEqual(match.orientation, "flipV")
        self.assertEqual(match.severity, "high")
        self.assertGreaterEqual(match.warp_ncc, 0.85)
        blank = np.full_like(original, 255)
        self.assertFalse(match_two_arrays(blank, blank).flagged)
        self.assertFalse(match_two_arrays(np.zeros_like(original), blank).flagged)

    def test_shared_axes_do_not_raise_high_and_reencoded_resize_remains_candidate(self):
        left = np.full((500, 500, 3), 255, dtype=np.uint8)
        right = left.copy()
        for image in (left, right):
            cv2.line(image, (50, 450), (450, 450), (0, 0, 0), 2)
            cv2.line(image, (50, 50), (50, 450), (0, 0, 0), 2)
        cv2.polylines(left, [np.array([[60, 400], [120, 300], [180, 350],
                                       [240, 150], [300, 270], [400, 130]])], False, (20, 20, 20), 3)
        cv2.polylines(right, [np.array([[60, 150], [120, 230], [180, 200],
                                        [240, 390], [300, 310], [400, 370]])], False, (20, 20, 20), 3)
        self.assertNotEqual(match_two_arrays(left, right).severity, "high")

        original = np.random.default_rng(31).integers(0, 255, (280, 320, 3), dtype=np.uint8)
        encoded = cv2.imencode(".jpg", cv2.resize(original, (448, 392)),
                               [cv2.IMWRITE_JPEG_QUALITY, 82])[1]
        match = match_two_arrays(original, cv2.imdecode(encoded, cv2.IMREAD_COLOR))
        self.assertTrue(match.flagged)
        self.assertGreaterEqual(match.warp_ncc, 0.85)
        self.assertNotEqual(match.severity, "high")

    def test_failed_step_does_not_count_as_complete(self):
        ocr = {"provider": "easyocr", "needed": 1, "done": 0, "failed": 1, "remaining": 0, "status": "failed"}
        companions = {key: [] for key in ("discovered", "scanned", "failed", "excluded")}
        numeric = {"detectors": {key: "not_evaluated" for key in ("B1", "B2", "B3", "B4", "B5")}}
        cnn = {"enabled": False, "status": "disabled"}
        stats = {"cross_image_pairs": {"possible": 3, "compared": 3, "verified": 0, "remaining": 0},
                 "panel_pairs": {"possible": 1, "compared": 0, "verified": 0, "remaining": 1}}
        coverage = coverage_from_steps([{"ok": False, "stats": stats}], None, ocr, companions, numeric, cnn, 3)
        self.assertEqual(coverage["image_pair"]["status"], "incomplete")
        self.assertEqual(coverage["panel_pair"]["remaining"], 1)
        self.assertEqual(coverage["whole_image"]["status"], "not_evaluated")

    def test_unicode_ocr_uses_pixels_and_retries_failure(self):
        class Reader:
            def readtext(self, pixels, **options):
                self.assert_pixels(pixels)
                return [([[1, 2], [20, 2], [20, 12], [1, 12]], "12.34", 0.98)]

            def assert_pixels(self, pixels):
                assert isinstance(pixels, np.ndarray) and pixels.shape == (48, 48, 3)

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = root / "中文图片.png"
            Image.fromarray(np.full((48, 48, 3), 120, dtype=np.uint8)).save(image)
            cache = root / "cache"
            cache.mkdir()
            digest = __import__("hashlib").sha256(image.read_bytes()).hexdigest()
            (cache / f"{digest}.json").write_text(json.dumps({"status": "failed", "reason": "old path decode"}), encoding="utf-8")
            with patch("manusift.local_ocr_runtime.get_reader", return_value=(Reader(), None)):
                result = scan([{"path": str(image), "image_id": "one"}], cache_dir=cache,
                              model_dir=root / "models", deadline=time.monotonic() + 30)
            self.assertEqual((result["done"], result["failed"]), (1, 0))
            self.assertEqual(result["records"][0]["numbers"][0]["text"], "12.34")

    def test_ocr_counts_cached_images_after_budget_expires(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = root / "cached.png"
            Image.fromarray(np.full((32, 32, 3), 90, dtype=np.uint8)).save(image)
            pending = root / "pending.png"
            Image.fromarray(np.full((32, 32, 3), 80, dtype=np.uint8)).save(pending)
            digest = __import__("hashlib").sha256(image.read_bytes()).hexdigest()
            cache = root / "cache"
            cache.mkdir()
            (cache / f"{digest}.json").write_text(json.dumps({"status": "done", "numbers": []}), encoding="utf-8")
            with patch("manusift.local_ocr_runtime.get_reader") as get_reader, patch("local_ocr.subprocess.run") as worker:
                result = scan_isolated([{"path": str(image)}, {"path": str(pending)}],
                                       cache_dir=cache, model_dir=root / "models", deadline=time.monotonic() - 1)
            get_reader.assert_not_called()
            worker.assert_not_called()
            self.assertEqual((result["needed"], result["done"], result["remaining"]), (2, 1, 1))

    def test_numeric_sources_and_coverage(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "source.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws["A1"], ws["D1"] = "Fig 3a", "Fig 4f"
            ws["B2"], ws["E2"] = 1.123456789, 1.123456789
            ws["A4"], ws["B4"] = "Histone code ChIPqPCR", "control"
            for row, value in ((5, 0.933209), (6, 1.019432), (7, 1.04736)):
                ws[f"B{row}"] = ws[f"H{row}"] = value
            ws["B10"] = 204.040404040404
            ws["A12"], ws["B12"] = "group one", 2
            ws["A13"], ws["B13"] = "group two", 3
            ws["A14"], ws["B14"] = "total", 6
            ws["C14"] = "=B12+B13"
            wb.save(path)
            result = audit([path])
            kinds = {item["detector"] for item in result["findings"]}
            self.assertTrue({"B1", "B2", "B4", "B5"} <= kinds)
            self.assertEqual(_calculate_formula("=B12+B13", {"B12": 2, "B13": 3}), 5)
            self.assertEqual(result["B5_recalculations"], 1)
            self.assertTrue(any("B4a" in x["recalculation"] for x in result["findings"]))
            self.assertTrue(any({s["figure"] for s in x["related_sources"]} == {"3A", "4F"}
                                for x in result["findings"] if x["detector"] == "B2"))

    def test_formula_index_is_shared_within_large_sheet(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "formulas.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            for row in range(1, 1001):
                sheet.cell(row, 1, row)
                sheet.cell(row, 2, f"=A{row}+1")
            workbook.save(path)
            values, formulas, errors = _xlsx_rows(path)
            self.assertFalse(errors)
            self.assertEqual((len(values), len(formulas)), (1000, 1000))
            self.assertTrue(all(item["cells"] is formulas[0]["cells"] for item in formulas))
            self.assertEqual(len(formulas[0]["cells"]), 2000)

    def test_cross_image_panel_queue_resumes(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            rng = np.random.default_rng(23)
            panel = rng.integers(0, 255, (180, 180, 3), dtype=np.uint8)
            images = []
            for index in range(2):
                canvas = np.full((180, 370, 3), 255, dtype=np.uint8)
                canvas[:, :180] = panel
                canvas[:, 190:] = rng.integers(0, 255, (180, 180, 3), dtype=np.uint8)
                path = root / f"figure-{index}.png"
                Image.fromarray(canvas).save(path)
                images.append(ExtractedImage(page=index, index=index, xref=index + 1, phash="",
                                             width=370, height=180, bytes_size=path.stat().st_size,
                                             image_path=str(path), exif={"zerowall_source": str(path)}))
            doc = ParsedDoc("test", str(root), [], images, {})
            settings = {"MANUSIFT_IMAGE_PAIR_CACHE_DIR": str(root / "cache"),
                        "MANUSIFT_CROSS_SIFT_RESUME": "1", "MANUSIFT_CROSS_SIFT_RUN_ID": "first",
                        "MANUSIFT_CROSS_SIFT_MAX_PAIRS": "1"}
            with patch.dict(os.environ, settings):
                _, first = _panel_then_match_findings(doc)
            self.assertGreater(first["possible"], 1)
            self.assertEqual(first["compared"], 1)
            with Path(first["remaining_pairs_path"]).open(encoding="utf-8") as pending:
                self.assertEqual(sum(1 for _ in pending), first["remaining"])
            settings["MANUSIFT_CROSS_SIFT_RUN_ID"] = "second"
            settings["MANUSIFT_CROSS_SIFT_MAX_PAIRS"] = "0"
            with patch.dict(os.environ, settings):
                findings, done = _panel_then_match_findings(doc)
            self.assertEqual(done["remaining"], 0)
            self.assertTrue(any(f.raw.get("image_a", {}).get("index") != f.raw.get("image_b", {}).get("index")
                                for f in findings))

    def test_cross_document_panel_pairs_are_not_skipped(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            rng = np.random.default_rng(28)
            images = []
            for index in range(2):
                canvas = rng.integers(0, 255, (180, 370, 3), dtype=np.uint8)
                path = root / f"paper-{index}.png"
                Image.fromarray(canvas).save(path)
                images.append(ExtractedImage(page=0, index=index, xref=1, phash="",
                    width=370, height=180, bytes_size=path.stat().st_size,
                    image_path=str(path), exif={"zerowall_source": str(path), "zerowall_document": index}))
            doc = ParsedDoc("test", str(root), [], images, {})
            settings = {"MANUSIFT_IMAGE_PAIR_CACHE_DIR": str(root / "cache"),
                        "MANUSIFT_CROSS_DOCUMENT_ONLY": "1",
                        "MANUSIFT_CROSS_SIFT_RUN_ID": "cross-document",
                        "MANUSIFT_CROSS_SIFT_MAX_PAIRS": "1"}
            with patch.dict(os.environ, settings):
                _, stats = _panel_then_match_findings(doc)
            self.assertGreater(stats["possible"], 0)
            self.assertEqual(stats["compared"], 1)
            self.assertEqual(stats["remaining"], stats["possible"] - 1)

    def test_panel_queue_does_not_retain_every_decoded_image(self):
        from manusift.detectors import sift_copymove

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            rng = np.random.default_rng(51)
            images = []
            for index in range(12):
                path = root / f"figure-{index}.png"
                Image.fromarray(rng.integers(0, 255, (180, 370, 3), dtype=np.uint8)).save(path)
                images.append(ExtractedImage(page=index, index=index, xref=index + 1,
                    phash="", width=370, height=180, bytes_size=path.stat().st_size,
                    image_path=str(path), exif={"zerowall_source": str(path)}))
            observed = []
            read_image = sift_copymove._read_image

            def tracked_read(path):
                array = read_image(path)
                observed.append(weakref.ref(array))
                return array

            def verify_pair(a, b):
                self.assertLessEqual(sum(ref() is not None for ref in observed), 2)
                return CrossMatchAnalysis(ok=True, reason="test")

            with patch.object(sift_copymove, "_read_image", side_effect=tracked_read), \
                 patch.object(sift_copymove, "match_two_arrays", side_effect=verify_pair), \
                 patch.dict(os.environ, {"MANUSIFT_CROSS_SIFT_MAX_PAIRS": "1",
                     "MANUSIFT_IMAGE_PAIR_CACHE_DIR": str(root / "cache"),
                     "MANUSIFT_CROSS_SIFT_RUN_ID": "bounded"}):
                _, stats = _panel_then_match_findings(ParsedDoc("bounded", str(root), [], images, {}))
            self.assertGreater(stats["possible"], 0)
            self.assertEqual(stats["compared"], 1)


if __name__ == "__main__":
    unittest.main()
