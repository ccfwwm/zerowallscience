#!/usr/bin/env python3
"""
patch_xlsx_jif.py  —  ZeroWall Literature helper
=================================================
Patch papers.xlsx in-place after `finalize`:
  - Fill 最新影响因子 and 影响因子年份 in 引文列表 where still empty,
    using verified_data.json JIF table.
  - Works with ANY task directory (not hardcoded to wang-2008).

Usage:
  python scripts/patch_xlsx_jif.py literature/<slug>
  # or from workspace root:
  python scripts/patch_xlsx_jif.py .   (if cwd IS the task dir)
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
from pathlib import Path

SCRIPT_DIR = Path(__file__).parent


def _load_jif_table(task_root: "Path | None" = None) -> dict[str, list]:
    """Two-level JIF table lookup: task-specific file wins over skill-level defaults."""
    def _load(path: Path) -> dict:
        if path.is_file():
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
                return data.get("jif", {}) if isinstance(data, dict) else {}
            except (OSError, ValueError):
                pass
        return {}

    skill_jif = _load(SCRIPT_DIR / "verified_data.json")
    task_jif: dict = {}
    if task_root is not None:
        task_jif = _load(Path(task_root) / "analysis" / "verified_data.json")

    # Task overrides skill; both contribute
    return {**skill_jif, **task_jif}


def _jif_key(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "").lower()
    s = re.sub(r"['\u2019\u2018\u00ad\-\u2013\u2014]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def match_jif(journal_name: str, jif_table: dict[str, list]):
    key = _jif_key(journal_name)
    hit = jif_table.get(key)
    if hit:
        return hit
    # partial prefix match on first 4 words
    words = key.split()[:4]
    if not words:
        return None
    prefix = " ".join(words)
    for k, v in jif_table.items():
        if k.startswith(prefix):
            return v
    return None


def main() -> int:
    try:
        import openpyxl
    except ImportError:
        print("openpyxl not installed — run: pip install openpyxl", file=sys.stderr)
        return 1

    task_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(".")
    if not task_dir.is_dir():
        task_dir = Path.cwd() / task_dir

    xlsx_path = task_dir / "papers.xlsx"
    if not xlsx_path.is_file():
        print(f"papers.xlsx not found in: {task_dir}", file=sys.stderr)
        return 1

    jif_table = _load_jif_table(task_dir)
    if not jif_table:
        print("No JIF table found in verified_data.json (task or skill level)", file=sys.stderr)
        return 1

    wb = openpyxl.load_workbook(str(xlsx_path))
    if "引文列表" not in wb.sheetnames:
        print("Sheet '引文列表' not found", file=sys.stderr)
        return 1

    ws = wb["引文列表"]
    header = [cell.value for cell in ws[1]]
    col_idx = {v: i + 1 for i, v in enumerate(header)}

    required = ("最新影响因子", "影响因子年份", "引用杂志全名")
    missing = [c for c in required if c not in col_idx]
    if missing:
        print(f"Missing columns in 引文列表: {missing}", file=sys.stderr)
        return 1

    jif_col    = col_idx["最新影响因子"]
    year_col   = col_idx["影响因子年份"]
    journal_col = col_idx["引用杂志全名"]

    patched = 0
    for row_idx in range(2, ws.max_row + 1):
        journal = ws.cell(row_idx, journal_col).value or ""
        if not journal or str(journal).strip().lower() == "pubmed":
            continue
        existing_jif = ws.cell(row_idx, jif_col).value
        if existing_jif not in (None, ""):
            continue
        result = match_jif(str(journal), jif_table)
        if not result:
            print(f"  no match: {journal}")
            continue
        jif_val, year_val, _ = result[0], result[1], result[2]
        ws.cell(row_idx, jif_col).value = jif_val
        ws.cell(row_idx, year_col).value = year_val
        patched += 1
        print(f"  patched: {journal} -> {jif_val} ({year_val})")

    print(f"\nTotal patched: {patched}")
    wb.save(str(xlsx_path))
    print(f"Saved {xlsx_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
